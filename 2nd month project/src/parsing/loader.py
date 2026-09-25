"""
Document Loader for PDF and DOCX files.
"""

import io
import re
import logging
from pathlib import Path
from typing import Dict, Any, Union, BinaryIO, List

logger = logging.getLogger(__name__)

try:
    import pypdf
except ImportError:
    pypdf = None

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None


def clean_text(raw_text: str) -> str:
    """
    Clean and normalize extracted text:
    - Replace Windows carriage returns
    - Remove non-printable control characters
    - Collapse excessive blank lines to max 2
    - Collapse horizontal multiple whitespace to single space
    """
    if not raw_text:
        return ""
    # Normalize carriage returns
    text = raw_text.replace("\r\n", "\n").replace("\r", "\n")
    # Remove control characters except standard tabs and newlines
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    # Collapse multiple horizontal spaces/tabs to a single space
    text = re.sub(r"[ \t]+", " ", text)
    # Collapse 3+ consecutive newlines to 2
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _extract_pdf_pages(stream_or_path: Union[BinaryIO, Path, bytes]) -> List[str]:
    """Extract text pages from PDF stream or path using available PDF libraries."""
    pages = []

    # Method 1: pypdf
    if pypdf is not None:
        try:
            if isinstance(stream_or_path, bytes):
                reader = pypdf.PdfReader(io.BytesIO(stream_or_path))
            elif isinstance(stream_or_path, Path):
                with open(stream_or_path, "rb") as f:
                    reader = pypdf.PdfReader(f)
                    return [page.extract_text() or "" for page in reader.pages]
            else:
                reader = pypdf.PdfReader(stream_or_path)
            return [page.extract_text() or "" for page in reader.pages]
        except Exception as e:
            logger.warning(f"pypdf extraction failed, trying fallbacks: {e}")

    # Method 2: PyPDF2
    if PyPDF2 is not None:
        try:
            if isinstance(stream_or_path, bytes):
                reader = PyPDF2.PdfReader(io.BytesIO(stream_or_path))
            elif isinstance(stream_or_path, Path):
                with open(stream_or_path, "rb") as f:
                    reader = PyPDF2.PdfReader(f)
                    return [page.extract_text() or "" for page in reader.pages]
            else:
                reader = PyPDF2.PdfReader(stream_or_path)
            return [page.extract_text() or "" for page in reader.pages]
        except Exception as e:
            logger.warning(f"PyPDF2 extraction failed, trying fallbacks: {e}")

    # Method 3: pdfplumber
    if pdfplumber is not None:
        try:
            if isinstance(stream_or_path, bytes):
                with pdfplumber.open(io.BytesIO(stream_or_path)) as pdf:
                    return [page.extract_text() or "" for page in pdf.pages]
            elif isinstance(stream_or_path, Path):
                with pdfplumber.open(stream_or_path) as pdf:
                    return [page.extract_text() or "" for page in pdf.pages]
            else:
                with pdfplumber.open(stream_or_path) as pdf:
                    return [page.extract_text() or "" for page in pdf.pages]
        except Exception as e:
            logger.warning(f"pdfplumber extraction failed, trying fallbacks: {e}")

    # Method 4: fitz (PyMuPDF)
    if fitz is not None:
        try:
            if isinstance(stream_or_path, bytes):
                doc = fitz.open(stream=stream_or_path, filetype="pdf")
            elif isinstance(stream_or_path, Path):
                doc = fitz.open(str(stream_or_path))
            else:
                doc = fitz.open(stream=stream_or_path.read(), filetype="pdf")
            return [page.get_text() or "" for page in doc]
        except Exception as e:
            logger.warning(f"PyMuPDF extraction failed: {e}")

    if not pages:
        raise ImportError(
            "No functional PDF parsing library found. Please install `pypdf` (`pip install pypdf`)."
        )
    return pages


def load_pdf(file_source: Union[str, Path, BinaryIO, bytes]) -> Dict[str, Any]:
    """
    Extract text and metadata from a PDF file or byte stream.
    """
    if isinstance(file_source, (str, Path)):
        path = Path(file_source)
        filename = path.name
        pages = _extract_pdf_pages(path)
    elif isinstance(file_source, bytes):
        filename = "uploaded_document.pdf"
        pages = _extract_pdf_pages(file_source)
    else:
        filename = getattr(file_source, "name", "uploaded_document.pdf")
        pages = _extract_pdf_pages(file_source)

    cleaned_pages = [clean_text(p) for p in pages]
    full_text = "\n\n--- Page Break ---\n\n".join([p for p in cleaned_pages if p.strip()])

    return {
        "filename": filename,
        "document_type": "pdf",
        "page_count": len(pages),
        "pages": cleaned_pages,
        "text": full_text or clean_text(" ".join(cleaned_pages)),
    }

def load_docx(file_source: Union[str, Path, BinaryIO, bytes]) -> Dict[str, Any]:
    """
    Extract text and metadata from a DOCX file or byte stream.
    """
    if isinstance(file_source, (str, Path)):
        path = Path(file_source)
        filename = path.name
        doc = docx.Document(path)
    elif isinstance(file_source, bytes):
        filename = "uploaded_document.docx"
        doc = docx.Document(io.BytesIO(file_source))
    else:
        filename = getattr(file_source, "name", "uploaded_document.docx")
        doc = docx.Document(file_source)

    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    
    # Also extract text from any tables
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            row_content = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
            if row_content:
                table_texts.append(row_content)

    all_content = "\n".join(paragraphs)
    if table_texts:
        all_content += "\n\n" + "\n".join(table_texts)

    cleaned_text = clean_text(all_content)

    return {
        "filename": filename,
        "document_type": "docx",
        "page_count": 1,
        "paragraphs_count": len(paragraphs),
        "text": cleaned_text,
    }

def load_document(file_source: Union[str, Path, BinaryIO, bytes], filename: str = "") -> Dict[str, Any]:
    """
    Unified loader detecting file type and extracting text.
    """
    detected_name = filename
    if not detected_name:
        if isinstance(file_source, (str, Path)):
            detected_name = str(file_source)
        else:
            detected_name = getattr(file_source, "name", "document.pdf")

    lower_name = detected_name.lower()
    if lower_name.endswith(".pdf"):
        res = load_pdf(file_source)
        if filename:
            res["filename"] = filename
        return res
    elif lower_name.endswith(".docx"):
        res = load_docx(file_source)
        if filename:
            res["filename"] = filename
        return res
    elif lower_name.endswith(".txt") or lower_name.endswith(".md"):
        if isinstance(file_source, (str, Path)):
            with open(file_source, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        elif isinstance(file_source, bytes):
            content = file_source.decode("utf-8", errors="replace")
        else:
            content = file_source.read().decode("utf-8", errors="replace")
        return {
            "filename": Path(detected_name).name,
            "document_type": "text",
            "page_count": 1,
            "text": clean_text(content),
        }
    else:
        raise ValueError(f"Unsupported document format: {detected_name}. Supported formats: PDF, DOCX, TXT, MD.")
