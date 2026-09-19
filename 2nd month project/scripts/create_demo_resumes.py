"""
Generates sample candidate resumes in PDF, DOCX, and TXT formats for demonstration.
"""

from pathlib import Path
import docx

PROJECT_ROOT = Path(__file__).resolve().parent.parent
RESUMES_DIR = PROJECT_ROOT / "data/resumes"
RESUMES_DIR.mkdir(parents=True, exist_ok=True)

RESUME_TEXT = """Alex Rivera
Senior Machine Learning & Generative AI Engineer
alex.rivera@example.com | (415) 555-0192 | San Francisco, CA | linkedin.com/in/alex-rivera-demo

PROFESSIONAL SUMMARY
Senior Machine Learning Engineer with 5+ years of experience designing, deploying, and optimizing end-to-end AI applications, RAG pipelines, and scalable microservices. Proven track record in vector search with FAISS, foundation model orchestration, and production FastAPI services.

TECHNICAL SKILLS
- Languages: Python, SQL, C++, TypeScript
- ML & GenAI: PyTorch, TensorFlow, Scikit-Learn, LangChain, FAISS, Transformers, RAG, Prompt Engineering
- Cloud & Data: AWS, Docker, Kubernetes, PostgreSQL, Redis, Pandas, NumPy
- Developer Tools: Git, Linux, CI/CD, FastAPI, Streamlit

PROFESSIONAL EXPERIENCE
Nexus AI Systems | Senior Generative AI Engineer | 2022 - Present
- Architected enterprise Retrieval-Augmented Generation (RAG) pipeline leveraging FAISS vector search and Gemini models, reducing question answering hallucinations by 42%.
- Deployed high-throughput FastAPI inference microservices processing over 1.5 million queries daily with sub-150ms p99 latency.
- Containerized model inference pipelines using Docker and Kubernetes on AWS ECS, implementing autoscaling based on GPU utilization.

Apex Analytics Corp | Machine Learning Engineer | 2019 - 2022
- Developed customer churn and risk prediction models in Scikit-Learn and PyTorch, increasing predictive accuracy by 18%.
- Optimized database queries and ETL pipelines in PostgreSQL and Pandas, accelerating weekly analytics training cycles by 3.5x.
- Created interactive internal monitoring dashboards in Streamlit to visualize model drift and prediction confidence intervals.

EDUCATION
University of California, Berkeley
Bachelor of Science in Computer Science | 2015 - 2019

NOTABLE PROJECTS
SmartHire RAG Career Engine
- Built open-source career matching portal combining FAISS semantic vector search with Gemini structured parsing.
- Implemented Human-in-the-Loop approval workflows and guardrail protection against prompt injection attacks.
"""

def create_docx():
    doc = docx.Document()
    doc.add_heading("Alex Rivera", level=1)
    doc.add_paragraph("Senior Machine Learning & Generative AI Engineer\nalex.rivera@example.com | (415) 555-0192 | San Francisco, CA")
    
    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph(
        "Senior Machine Learning Engineer with 5+ years of experience designing, deploying, and optimizing end-to-end AI applications, RAG pipelines, and scalable microservices."
    )

    doc.add_heading("Technical Skills", level=2)
    doc.add_paragraph("Python, SQL, PyTorch, Scikit-Learn, LangChain, FAISS, Docker, Kubernetes, AWS, PostgreSQL, FastAPI, Streamlit")

    doc.add_heading("Professional Experience", level=2)
    p1 = doc.add_paragraph()
    p1.add_run("Senior Generative AI Engineer - Nexus AI Systems (2022 - Present)\n").bold = True
    p1.add_run("- Architected enterprise RAG pipeline using FAISS and Gemini, reducing hallucinations by 42%.\n")
    p1.add_run("- Deployed FastAPI microservices serving 1.5M requests daily with sub-150ms latency.\n")
    p1.add_run("- Implemented containerized deployment using Docker and Kubernetes on AWS.")

    p2 = doc.add_paragraph()
    p2.add_run("Machine Learning Engineer - Apex Analytics Corp (2019 - 2022)\n").bold = True
    p2.add_run("- Developed risk prediction models in Scikit-Learn and PyTorch, improving accuracy by 18%.\n")
    p2.add_run("- Optimized database queries in PostgreSQL, accelerating training pipelines by 3.5x.")

    doc.add_heading("Education", level=2)
    doc.add_paragraph("University of California, Berkeley\nBachelor of Science in Computer Science (2015 - 2019)")

    out_path = RESUMES_DIR / "sample_resume.docx"
    doc.save(str(out_path))
    print(f"Created DOCX: {out_path}")

def escape_pdf(t: str) -> str:
    return t.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

def page_stream(lines: list) -> bytes:
    cmds = ["BT", "/F1 10 Tf", "54 750 Td", "13 TL"]
    for l in lines:
        cmds.append(f"({escape_pdf(l)}) Tj")
        cmds.append("T*")
    cmds.append("ET")
    return "\n".join(cmds).encode("latin-1")

def create_pdf():
    lines = [l for l in RESUME_TEXT.splitlines() if l.strip()]
    stream = page_stream(lines[:45])
    
    objects = []
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    objects.append(b"<< /Type /Pages /Kids [4 0 R] /Count 1 >>")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    objects.append(b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 3 0 R >> >> /Contents 5 0 R >>")
    objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")

    output = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")
    offsets = [0]
    for idx, obj in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(f"{idx} 0 obj\n".encode())
        output.extend(obj)
        output.extend(b"\nendobj\n")

    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")
    for off in offsets[1:]:
        output.extend(f"{off:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())

    out_path = RESUMES_DIR / "sample_resume.pdf"
    out_path.write_bytes(bytes(output))
    print(f"Created PDF: {out_path}")

    txt_path = RESUMES_DIR / "sample_resume.txt"
    txt_path.write_text(RESUME_TEXT, encoding="utf-8")
    print(f"Created TXT: {txt_path}")

if __name__ == "__main__":
    create_docx()
    create_pdf()
